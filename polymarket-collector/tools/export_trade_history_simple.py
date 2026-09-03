#!/usr/bin/env python3
"""Read-only export of the Polymarket fill history as one clean 'Trade History' sheet.

Built for manual verification against the Polymarket website History tab, so the
layout is deliberately minimal: one row per executed fill, twelve columns, no
summaries, no JSON, no derived analytics.

Read-only: HTTP GET only. No orders, no database writes, no service interaction.
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import hmac
import os
import sys
import time
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

JERUSALEM = ZoneInfo("Asia/Jerusalem")
DEFAULT_ENV_FILE = "/etc/polymarket-live/trader.env"
DEFAULT_OUTPUT_DIR = "/opt/polymarket-btc-live/output"

COLUMNS = [
    "Date & Time Jerusalem",
    "Action",
    "Outcome",
    "Price",
    "Tokens",
    "Trade Amount USD",
    "Maker / Taker",
    "Fee USD",
    "Fee Source",
    "Event Slug",
    "Transaction Hash",
    "Trade ID",
]


def dec(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return result if result.is_finite() else None


def load_env_file(path: str) -> dict[str, str]:
    values: dict[str, str] = {}
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def resolve_config(env_file: str) -> dict[str, str]:
    try:
        file_values = load_env_file(env_file)
    except PermissionError:
        raise SystemExit(f"cannot read {env_file}; re-run with sudo")
    except FileNotFoundError:
        file_values = {}

    def pick(name: str, default: str = "") -> str:
        return os.environ.get(name) or file_values.get(name) or default

    config = {
        "api_key": pick("POLYMARKET_API_KEY"),
        "api_secret": pick("POLYMARKET_API_SECRET"),
        "api_passphrase": pick("POLYMARKET_API_PASSPHRASE"),
        "signer_address": pick("POLYMARKET_SIGNER_ADDRESS"),
        "proxy_wallet": pick("POLYMARKET_FUNDER_ADDRESS", pick("POLYMARKET_PROFILE_ADDRESS")),
        "clob_host": pick("POLYMARKET_CLOB_HOST", "https://clob.polymarket.com").rstrip("/"),
        "data_host": pick("POLYMARKET_DATA_API_HOST", "https://data-api.polymarket.com").rstrip("/"),
    }
    missing = [k for k, v in config.items() if not v]
    if missing:
        raise SystemExit(f"missing configuration: {', '.join(missing)}")
    return config


def fetch_clob_trades(config: dict[str, str]) -> list[dict[str, Any]]:
    path = "/data/trades"
    trades: list[dict[str, Any]] = []
    cursor: str | None = None
    with httpx.Client(timeout=45) as client:
        while True:
            timestamp = int(time.time())
            secret = config["api_secret"]
            raw = base64.urlsafe_b64decode(secret + "=" * ((-len(secret)) % 4))
            digest = hmac.new(raw, f"{timestamp}GET{path}".encode(), hashlib.sha256).digest()
            response = client.get(
                config["clob_host"] + path,
                params={"next_cursor": cursor} if cursor else {},
                headers={
                    "POLY_ADDRESS": config["signer_address"],
                    "POLY_API_KEY": config["api_key"],
                    "POLY_PASSPHRASE": config["api_passphrase"],
                    "POLY_SIGNATURE": base64.urlsafe_b64encode(digest).decode("ascii"),
                    "POLY_TIMESTAMP": str(timestamp),
                },
            )
            response.raise_for_status()
            payload = response.json()
            trades.extend(payload.get("data") or [])
            nxt = payload.get("next_cursor")
            if not nxt or nxt == "LTE=" or nxt == cursor:
                break
            cursor = nxt
    return trades


def fetch_data_trades(config: dict[str, str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset, size = 0, 500
    with httpx.Client(timeout=45) as client:
        while True:
            response = client.get(
                config["data_host"] + "/trades",
                params={"user": config["proxy_wallet"], "takerOnly": "false",
                        "limit": size, "offset": offset},
            )
            response.raise_for_status()
            page = response.json()
            rows.extend(page)
            if len(page) < size:
                break
            offset += len(page)
    return rows


def build_rows(clob: list[dict[str, Any]], data: list[dict[str, Any]],
               proxy: str) -> list[dict[str, Any]]:
    """One row per fill, always from our account's own perspective.

    The CLOB top-level side/price/asset describe the TAKER of the match, so when
    we were the MAKER those fields belong to the counterparty. In that case the
    account's real execution is rebuilt from the maker_orders entries owned by
    our proxy wallet.
    """
    by_tx: dict[str, list[dict[str, Any]]] = {}
    for row in data:
        by_tx.setdefault(str(row.get("transactionHash", "")).lower(), []).append(row)

    rows: list[dict[str, Any]] = []
    for trade in clob:
        trader_side = str(trade.get("trader_side") or "").upper()
        mine = [
            order for order in (trade.get("maker_orders") or [])
            if isinstance(order, dict)
            and str(order.get("maker_address", "")).lower() == proxy
        ]

        if trader_side == "MAKER" and mine:
            size = sum((dec(o.get("matched_amount")) or Decimal(0) for o in mine), Decimal(0))
            weighted = sum(
                ((dec(o.get("price")) or Decimal(0)) * (dec(o.get("matched_amount")) or Decimal(0))
                 for o in mine), Decimal(0))
            price = (weighted / size) if size > 0 else dec(mine[0].get("price"))
            side = str(mine[0].get("side") or "").upper()
            outcome = str(mine[0].get("outcome") or "")
        else:
            size, price = dec(trade.get("size")), dec(trade.get("price"))
            side = str(trade.get("side") or "").upper()
            outcome = str(trade.get("outcome") or "")

        tx_hash = str(trade.get("transaction_hash") or "")
        candidates = by_tx.get(tx_hash.lower(), [])
        match = candidates[0] if len(candidates) == 1 else None
        if match is None and len(candidates) > 1:
            exact = [c for c in candidates
                     if str(c.get("side", "")).upper() == side and dec(c.get("size")) == size]
            match = exact[0] if len(exact) == 1 else None

        # Prefer the Data API timestamp so rows line up with the website History.
        epoch = None
        if match and match.get("timestamp"):
            epoch = int(dec(match.get("timestamp")) or 0)
        if not epoch:
            epoch = int(dec(trade.get("match_time")) or 0)
        local = datetime.fromtimestamp(epoch, tz=timezone.utc).astimezone(JERUSALEM)

        # Polymarket returns a fee RATE, never a settled fee amount.
        rate_bps = dec(trade.get("fee_rate_bps"))
        if rate_bps is None or price is None or size is None:
            fee, fee_source = None, "Unavailable"
        elif rate_bps == 0:
            fee, fee_source = 0.0, "Calculated"
        else:
            fee = float((size * price * (Decimal(1) - price) * rate_bps / Decimal(10000))
                        .quantize(Decimal("0.000001")))
            fee_source = "Calculated"

        rows.append({
            "Date & Time Jerusalem": local.strftime("%d/%m/%Y %H:%M:%S"),
            "Action": "קניה" if side == "BUY" else "מכירה" if side == "SELL" else side,
            "Outcome": (match.get("outcome") if match else outcome) or outcome,
            "Price": float(price) if price is not None else None,
            "Tokens": float(size) if size is not None else None,
            "Trade Amount USD": (float(price * size)
                                 if price is not None and size is not None else None),
            "Maker / Taker": trader_side,
            "Fee USD": fee,
            "Fee Source": fee_source,
            "Event Slug": (match.get("eventSlug") if match else "") or "",
            "Transaction Hash": tx_hash,
            "Trade ID": trade.get("id"),
            "_sort": epoch,
        })

    rows.sort(key=lambda r: r["_sort"])
    for row in rows:
        row.pop("_sort")
    return rows


def write_workbook(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trade History"

    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append([row[column] for column in COLUMNS])

    widths = {"Date & Time Jerusalem": 21, "Action": 10, "Outcome": 12, "Price": 14,
              "Tokens": 14, "Trade Amount USD": 18, "Maker / Taker": 14, "Fee USD": 12,
              "Fee Source": 13, "Event Slug": 26, "Transaction Hash": 68, "Trade ID": 38}
    formats = {"Price": "0.00########", "Tokens": "0.00######",
               "Trade Amount USD": "0.000000", "Fee USD": "0.000000"}
    for index, column in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = widths[column]
        if column in formats:
            for cell in sheet[letter][1:]:
                cell.number_format = formats[column]
        elif column in ("Action", "Outcome", "Maker / Taker", "Fee Source"):
            for cell in sheet[letter][1:]:
                cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def assert_no_secrets(path: Path, config: dict[str, str]) -> None:
    import zipfile

    blob = b""
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            blob += archive.read(name)
    for label in ("api_key", "api_secret", "api_passphrase"):
        if config[label].encode() in blob:
            path.unlink(missing_ok=True)
            raise SystemExit(f"ABORTED: {label} leaked into the workbook; file deleted")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--env-file", default=DEFAULT_ENV_FILE)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()

    config = resolve_config(args.env_file)
    clob = fetch_clob_trades(config)
    data = fetch_data_trades(config)
    print(f"[info] CLOB fills: {len(clob)} | Data API rows: {len(data)}")

    rows = build_rows(clob, data, config["proxy_wallet"].lower())

    started = datetime.now(tz=JERUSALEM)
    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"polymarket_trade_history_{started:%Y%m%d_%H%M%S}.xlsx"
    write_workbook(rows, path)
    assert_no_secrets(path, config)

    uid, gid = os.environ.get("SUDO_UID"), os.environ.get("SUDO_GID")
    if uid and gid:
        os.chown(path, int(uid), int(gid))

    buys = sum(1 for r in rows if r["Action"] == "קניה")
    sells = sum(1 for r in rows if r["Action"] == "מכירה")
    makers = sum(1 for r in rows if r["Maker / Taker"] == "MAKER")
    takers = sum(1 for r in rows if r["Maker / Taker"] == "TAKER")
    print(f"[ok] {path}")
    print(f"[ok] rows={len(rows)} first={rows[0]['Date & Time Jerusalem'] if rows else '-'} "
          f"last={rows[-1]['Date & Time Jerusalem'] if rows else '-'}")
    print(f"[ok] buy={buys} sell={sells} maker={makers} taker={takers}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
